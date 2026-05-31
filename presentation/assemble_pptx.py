#!/usr/bin/env python3
"""Assemble the rendered TED slides into a full-bleed 16:9 .pptx."""
import os, glob
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
SL = os.path.join(HERE, "slides")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

files = sorted(glob.glob(os.path.join(SL, "*.png")))
for f in files:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(f, 0, 0, width=prs.slide_width, height=prs.slide_height)

out = os.path.join(HERE, "Liver_Support_ICU_TED_Keynote.pptx")
prs.save(out)
print("SAVED", out, "—", len(prs.slides._sldIdLst), "slides,",
      round(os.path.getsize(out)/1024/1024, 1), "MB")
